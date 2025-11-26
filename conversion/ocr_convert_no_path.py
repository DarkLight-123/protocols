import os
import sys
import time
import argparse
from pathlib import Path

import fitz  # PyMuPDF
from PIL import Image, ImageEnhance
import pytesseract

import cv2
import numpy as np
from docx import Document

# --- Локальный путь к языковым файлам (без PATH и без прав админа) ---
# rus.traineddata должен лежать внутри этой папки
LOCAL_TESSDATA_DIR = r"C:\Users\ondre\tessdata"

# --- Путь к tesseract.exe (без PATH) ---
DEFAULT_TESS_PATHS = [
    r"C:\Program Files\Tesseract-OCR\tesseract.exe",
    r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
]


def resolve_tesseract_path():
    for p in DEFAULT_TESS_PATHS:
        if os.path.isfile(p):
            return p
    # Если установлен куда-то ещё — пропиши путь вручную:
    # return r"C:\...\tesseract.exe"
    return None


def preprocess_for_ocr(pil_img: Image.Image) -> Image.Image:
    """Агрессивная предобработка: обрезка полей, бинаризация, deskew, шарп."""
    gray = pil_img.convert("L")
    gray = ImageEnhance.Contrast(gray).enhance(1.6)
    gray = ImageEnhance.Brightness(gray).enhance(1.1)
    img = np.array(gray)

    # Обрезка белых полей
    t = cv2.threshold(img, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)[1]
    coords = cv2.findNonZero(t)
    if coords is not None:
        x, y, w, h = cv2.boundingRect(coords)
        img = img[y:y + h, x:x + w]

    # Шумоподавление + адаптивная бинаризация
    img = cv2.fastNlMeansDenoising(img, None, 15, 7, 21)
    bw = cv2.adaptiveThreshold(
        img, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
        cv2.THRESH_BINARY, 35, 15
    )

    # Deskew
    coords = np.column_stack(np.where(bw == 0))
    if coords.size > 0:
        rect = cv2.minAreaRect(coords)
        angle = rect[-1]
        angle = -(90 + angle) if angle < -45 else -angle
        h, w = bw.shape[:2]
        M = cv2.getRotationMatrix2D((w // 2, h // 2), angle, 1.0)
        bw = cv2.warpAffine(
            bw, M, (w, h),
            flags=cv2.INTER_CUBIC, borderMode=cv2.BORDER_REPLICATE
        )

    # Unsharp + лёгкая морфология
    blur = cv2.GaussianBlur(bw, (0, 0), 1.0)
    sharp = cv2.addWeighted(bw, 1.5, blur, -0.5, 0)
    kernel = np.ones((2, 2), np.uint8)
    sharp = cv2.morphologyEx(sharp, cv2.MORPH_OPEN, kernel, iterations=1)

    return Image.fromarray(sharp).convert("1")


def page_to_image(page, dpi=500):
    """Рендер страницы PDF в PIL.Image с высоким DPI."""
    zoom = dpi / 72.0
    mat = fitz.Matrix(zoom, zoom)
    pix = page.get_pixmap(matrix=mat, alpha=False)
    img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
    img.info["dpi"] = (dpi, dpi)
    return img


def ocr_image_to_text(image: Image.Image, lang: str) -> str:
    """OCR одной страницы с правильным конфигом и локальным tessdata."""
    bw = preprocess_for_ocr(image)
    config = rf'--oem 1 --psm 4 -c preserve_interword_spaces=1 --tessdata-dir "{LOCAL_TESSDATA_DIR}"'
    return pytesseract.image_to_string(bw, lang=lang, config=config)


LAT2CYR = str.maketrans({
    "A": "А", "B": "В", "E": "Е", "K": "К", "M": "М", "H": "Н", "O": "О",
    "P": "Р", "C": "С", "T": "Т", "X": "Х", "Y": "У",
    "a": "а", "e": "е", "o": "о", "p": "р", "c": "с", "x": "х", "y": "у",
    "k": "к", "m": "м", "h": "н", "t": "т"
})


def normalize_cyr(text: str) -> str:
    t = text.translate(LAT2CYR)
    t = (t.replace("0С", "ОС").replace("0с", "Ос")
         .replace("3а", "За").replace("3e", "Эе")
         .replace(" H ", " Н ").replace(" P ", " Р ").replace(" C ", " С ").replace(" T ", " Т "))
    return t


def ocr_by_blocks(image: Image.Image, lang: str) -> str:
    """Распознаёт по блокам (TSV), сортирует и склеивает строки."""
    bw = preprocess_for_ocr(image)
    config = rf'--oem 1 --psm 4 -c preserve_interword_spaces=1 --tessdata-dir "{LOCAL_TESSDATA_DIR}"'
    data = pytesseract.image_to_data(
        bw, lang=lang, config=config, output_type=pytesseract.Output.DICT
    )

    lines = []
    n = len(data["text"])
    for i in range(n):
        try:
            if int(data["conf"][i]) > 60 and data["text"][i].strip():
                lines.append((
                    data["page_num"][i], data["block_num"][i], data["par_num"][i], data["line_num"][i],
                    data["left"][i], data["top"][i], data["text"][i]
                ))
        except Exception:
            continue

    lines.sort(key=lambda x: (x[0], x[1], x[2], x[3], x[4]))

    out, buff, last_key = [], [], None
    for rec in lines:
        key = (rec[0], rec[1], rec[2], rec[3])
        if key != last_key and buff:
            out.append(" ".join(buff))
            buff = []
        buff.append(rec[6])
        last_key = key
    if buff:
        out.append(" ".join(buff))

    return normalize_cyr("\n".join(out))


def pick_langs(tess_cmd: str) -> str:
    """Выбираем язык: если rus.traineddata есть локально -> 'rus', иначе пытаемся определить, иначе 'eng'."""
    rus_local = os.path.isfile(os.path.join(LOCAL_TESSDATA_DIR, "rus.traineddata"))
    if rus_local:
        return "rus"

    # Попробуем спросить у tesseract список языков
    try:
        import subprocess
        proc = subprocess.run([tess_cmd, "--list-langs"], capture_output=True, text=True, timeout=5)
        langs_out = (proc.stdout + proc.stderr).lower()
        return "rus" if "rus" in langs_out else "eng"
    except Exception:
        return "eng"


def convert_pdf_to_docx(pdf_path: Path, out_dir: Path, lang: str) -> Path:
    """Конвертирует один PDF (сканы) в DOCX с распознанным текстом."""
    t0 = time.time()
    out_dir.mkdir(parents=True, exist_ok=True)
    docx_path = out_dir / (pdf_path.stem + ".docx")

    with fitz.open(pdf_path) as doc:
        word_doc = Document()
        for i, page in enumerate(doc, 1):
            img = page_to_image(page, dpi=500)  # высокий DPI
            # можно попробовать блоковый OCR, если обычный даёт крокозябры:
            # text = ocr_by_blocks(img, lang).strip()
            text = ocr_image_to_text(img, lang).strip()
            text = normalize_cyr(text)

            word_doc.add_paragraph(f"[Страница {i}]")
            if text:
                word_doc.add_paragraph(text)
            else:
                word_doc.add_paragraph("(Пусто / не распознано)")
            word_doc.add_page_break()

    word_doc.save(docx_path)
    print(f"[OK] {pdf_path.name} → {docx_path.name}  ({time.time() - t0:.1f}s)")
    return docx_path


def main():
    parser = argparse.ArgumentParser(
        description="OCR-конвертация PDF (сканы) → DOCX без PATH для Tesseract"
    )
    parser.add_argument("input", help="Путь к PDF файлу или папке с PDF")
    parser.add_argument("--out", default="word_files", help="Папка для сохранения DOCX (по умолчанию: word_files)")
    args = parser.parse_args()

    # Путь к tesseract.exe
    tess_cmd = resolve_tesseract_path()
    if not tess_cmd:
        print("❌ Не найден tesseract.exe.\n"
              "Укажи вручную путь в resolve_tesseract_path() в этом файле.\n"
              r'Пример: C:\Program Files\Tesseract-OCR\tesseract.exe')
        sys.exit(1)
    pytesseract.pytesseract.tesseract_cmd = tess_cmd

    lang = pick_langs(tess_cmd)
    print(f"Tesseract: {tess_cmd}")
    print(f"Язык OCR: {lang}")
    print(f"Tessdata dir: {LOCAL_TESSDATA_DIR}")

    inp = Path(args.input)
    out_dir = Path(args.out)

    if inp.is_file() and inp.suffix.lower() == ".pdf":
        convert_pdf_to_docx(inp, out_dir, lang)
    elif inp.is_dir():
        pdfs = sorted([p for p in inp.glob("*.pdf") if p.is_file()])
        if not pdfs:
            print("❗ В папке нет PDF.")
            sys.exit(2)
        for pdf in pdfs:
            try:
                convert_pdf_to_docx(pdf, out_dir, lang)
            except Exception as e:
                print(f"[FAIL] {pdf.name}: {e}")
    else:
        print("❌ Укажи путь к PDF файлу или папке с PDF.")
        sys.exit(3)


if __name__ == "__main__":
    main()
