# conversion/ocr_paddle.py
import os
import time
import argparse
from pathlib import Path

import fitz  # PyMuPDF
import numpy as np
import cv2
from PIL import Image
from docx import Document
from paddleocr import PaddleOCR

# --------- окружение/модели/потоки ----------
PADDLE_MODELS_DIR = r"C:\Users\ondre\OneDrive\Desktop\protocols\.paddleocr"
os.environ.setdefault("PADDLEOCR_HOME", PADDLE_MODELS_DIR)
os.environ.setdefault("OMP_NUM_THREADS", "4")   # ограничим потоки — на CPU так быстрее и стабильнее
os.environ.setdefault("MKL_NUM_THREADS", "4")

# --------- рендер PDF страницы в PIL.Image ----------
def page_to_image(page, dpi=240) -> Image.Image:
    zoom = dpi / 72.0
    mat = fitz.Matrix(zoom, zoom)
    pix = page.get_pixmap(matrix=mat, alpha=False)
    img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
    img.info["dpi"] = (dpi, dpi)
    return img

# --------- приведение размера (ускоряет детекцию) ----------
def limit_long_side(img_bgr: np.ndarray, max_side: int = 1600) -> np.ndarray:
    h, w = img_bgr.shape[:2]
    m = max(h, w)
    if m <= max_side:
        return img_bgr
    scale = max_side / float(m)
    new_w, new_h = int(w * scale), int(h * scale)
    return cv2.resize(img_bgr, (new_w, new_h), interpolation=cv2.INTER_AREA)

# --------- детект «синего фона» и его подавление ----------
def remove_blue_background(bgr: np.ndarray) -> np.ndarray:
    # HSV: синее ~ (h ∈ [90..140] при OpenCV [0..179])
    hsv = cv2.cvtColor(bgr, cv2.COLOR_BGR2HSV)
    lower = np.array([90,  30,  30], dtype=np.uint8)
    upper = np.array([140, 255, 255], dtype=np.uint8)
    mask_blue = cv2.inRange(hsv, lower, upper)

    # Ослабим синий: там, где маска, повышаем яркость/делаем белее
    inv = cv2.bitwise_not(mask_blue)
    clean = bgr.copy()
    # осветлим синие области, смешивая с белым
    white = np.full_like(bgr, 255)
    clean = cv2.addWeighted(clean, 1.0, white, 0.35, 0)  # слегка белим всё
    # но текст (в основном не синий) постараемся сохранить: вернём пиксели вне синей маски
    clean[inv > 0] = bgr[inv > 0]
    return clean

# --------- усиление/бинаризация/дескью ----------
def preprocess_for_ocr(bgr: np.ndarray) -> np.ndarray:
    # 1) уберём синий фон
    bgr = remove_blue_background(bgr)

    # 2) переведём в серый и применим CLAHE (локальный контраст)
    gray = cv2.cvtColor(bgr, cv2.COLOR_BGR2GRAY)
    clahe = cv2.createCLAHE(clipLimit=2.5, tileGridSize=(8, 8))
    gray = clahe.apply(gray)

    # 3) лёгкое шумоподавление
    gray = cv2.fastNlMeansDenoising(gray, None, 12, 7, 21)

    # 4) адаптивная бинаризация
    bw = cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
                               cv2.THRESH_BINARY, 35, 15)

    # 5) дескью по маске текста
    coords = np.column_stack(np.where(bw == 0))
    if coords.size > 0:
        rect = cv2.minAreaRect(coords)
        angle = rect[-1]
        angle = -(90 + angle) if angle < -45 else -angle
        h, w = bw.shape[:2]
        M = cv2.getRotationMatrix2D((w//2, h//2), angle, 1.0)
        bw = cv2.warpAffine(bw, M, (w, h), flags=cv2.INTER_CUBIC, borderMode=cv2.BORDER_REPLICATE)

    # 6) лёгкая морфология (снять мусор, подчеркнуть штрихи)
    kernel = np.ones((2, 2), np.uint8)
    bw = cv2.morphologyEx(bw, cv2.MORPH_OPEN, kernel, iterations=1)

    # возвратим 3-канальный для OCR
    return cv2.cvtColor(bw, cv2.COLOR_GRAY2BGR)

# --------- разбор результата нового/старого API ----------
def coerce_items(result_obj):
    items = []

    # Новый API: OCRResult (у элемента есть .texts, .boxes)
    try:
        results = result_obj if isinstance(result_obj, (list, tuple)) else [result_obj]
        handled = False
        for r in results:
            boxes = getattr(r, "boxes", None)
            texts = getattr(r, "texts", None)
            if boxes is not None and texts is not None:
                for txt, box in zip(texts, boxes):
                    if not txt:
                        continue
                    xs = [p[0] for p in box]
                    ys = [p[1] for p in box]
                    items.append((int(min(ys)), int(min(xs)), str(txt).strip()))
                handled = True
        if handled:
            return items
    except Exception:
        pass

    # Старый API: [ [box], (text, conf) ]
    if isinstance(result_obj, (list, tuple)) and result_obj:
        page = result_obj[0]
        if isinstance(page, (list, tuple)):
            for elem in page:
                if isinstance(elem, (list, tuple)) and len(elem) == 2:
                    box, txt_conf = elem
                    try:
                        txt = txt_conf[0]
                        xs = [p[0] for p in box]
                        ys = [p[1] for p in box]
                        items.append((int(min(ys)), int(min(xs)), str(txt).strip()))
                    except Exception:
                        pass

    return items

# --------- один прогон OCR с запасными вариантами ----------
def run_ocr_once(ocr: PaddleOCR, bgr: np.ndarray):
    # 1) предобработка
    prep = preprocess_for_ocr(limit_long_side(bgr))
    # 2) попытка на подготовленном
    try:
        r = ocr.predict(prep)
    except TypeError:
        r = ocr.ocr(prep)
    items = coerce_items(r)
    if len(items) > 0:
        return items, "prep"

    # 3) fallback: без жёсткой бинаризации (вдруг сработает лучше)
    try:
        r = ocr.predict(limit_long_side(bgr))
    except TypeError:
        r = ocr.ocr(limit_long_side(bgr))
    items = coerce_items(r)
    if len(items) > 0:
        return items, "orig"

    # 4) fallback: поворот на 180°
    bgr180 = cv2.rotate(bgr, cv2.ROTATE_180)
    try:
        r = ocr.predict(preprocess_for_ocr(limit_long_side(bgr180)))
    except TypeError:
        r = ocr.ocr(preprocess_for_ocr(limit_long_side(bgr180)))
    items = coerce_items(r)
    if len(items) > 0:
        return items, "prep_180"

    return [], "none"

def items_to_text(items):
    items.sort()  # по Y, затем по X
    lines = []
    cur = []
    last_y = None
    for y, x, t in items:
        if not t:
            continue
        if last_y is None or abs(y - last_y) < 20:
            cur.append(t)
        else:
            if cur:
                lines.append(" ".join(cur))
            cur = [t]
        last_y = y
    if cur:
        lines.append(" ".join(cur))
    return "\n".join(lines).strip()

# --------- основной цикл по PDF ----------
def convert_pdf_to_docx(pdf_path: Path, out_dir: Path, ocr: PaddleOCR) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    docx_path = out_dir / (pdf_path.stem + ".docx")
    word_doc = Document()

    print(f"\n=== Обрабатываю: {pdf_path.name} ===")
    t_file = time.time()

    with fitz.open(pdf_path) as doc:
        for i, page in enumerate(doc, 1):
            t0 = time.time()
            print(f"  - Страница {i}/{len(doc)} ...", end="", flush=True)

            pil = page_to_image(page, dpi=240)
            bgr = cv2.cvtColor(np.array(pil), cv2.COLOR_RGB2BGR)

            items, mode = run_ocr_once(ocr, bgr)
            text = items_to_text(items)

            word_doc.add_paragraph(f"[Страница {i}] (mode={mode}, blocks={len(items)})")
            word_doc.add_paragraph(text if text else "(Пусто / не распознано)")
            word_doc.add_page_break()

            print(f" {mode}, блоков: {len(items)}, {time.time()-t0:.1f} c")

    word_doc.save(docx_path)
    print(f"[OK] {pdf_path.name} → {docx_path.name} за {time.time()-t_file:.1f} c")
    return docx_path

def main():
    ap = argparse.ArgumentParser(description="OCR PDF → DOCX через PaddleOCR (русский, быстрый)")
    ap.add_argument("input", help="Путь к PDF файлу или папке")
    ap.add_argument("--out", default="word_files", help="Куда класть DOCX")
    args = ap.parse_args()

    # Новые параметры PaddleOCR без deprecated и с лимитами
    ocr = PaddleOCR(
        lang="ru",
        use_textline_orientation=False,        # быстрее; для протоколов обычно не нужно
        text_det_limit_side_len=1600,          # ограничение детектора по размеру
        text_recognition_batch_size=16,        # батч на распознавание
        cpu_threads=4                          # использовать 4 потока
    )

    inp = Path(args.input)
    out_dir = Path(args.out)

    if inp.is_file() and inp.suffix.lower() == ".pdf":
        convert_pdf_to_docx(inp, out_dir, ocr)
    elif inp.is_dir():
        pdfs = sorted([p for p in inp.glob("*.pdf") if p.is_file()])
        if not pdfs:
            print("❗ В папке нет PDF.")
            return
        for pdf in pdfs:
            try:
                convert_pdf_to_docx(pdf, out_dir, ocr)
            except Exception as e:
                print(f"[FAIL] {pdf.name}: {e}")
    else:
        print("❌ Укажи существующий PDF или папку с PDF.")

if __name__ == "__main__":
    main()


